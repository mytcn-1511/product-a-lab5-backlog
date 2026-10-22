# Ghi SWP và giờ ước lượng lên GitHub Project từ báo cáo HTML của skill estimate-swp và bộ ước tay.
# - SWP: lấy finalSWP (issue lá) và finalSWP đã cộng sẵn (issue cha) từ dữ liệu nhúng trong HTML; không tự tính lại.
# - Giờ: issue ước tay giữ số tay; issue cha ước tay chia giờ cho issue con theo tỷ lệ SWP;
#   issue còn lại = SWP × k (k = tổng giờ tay ÷ tổng SWP cùng bộ), làm tròn 0,5; issue cha = tổng con.
# Chạy lại an toàn: chỉ ghi ô có giá trị khác. Cách chạy:
#   python3 skills/ghi-swp-len-project.py --dry-run
#   python3 skills/ghi-swp-len-project.py
import argparse, json, re, subprocess, sys, time
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
HTML = sorted((ROOT / "swp-report").glob("estimate-swp-*.html"))[-1]
HAND = ROOT / "uoc-gio-tay" / "uoc-gio-tay-15-issue-20261021_filled.xlsx"
OWNER, NUMBER = "mytcn-1511", "1"


def gh(*args):
    r = subprocess.run(["gh", *args], capture_output=True, text=True)
    if r.returncode != 0:
        sys.exit(f"gh {' '.join(args[:3])} lỗi: {r.stderr.strip()}")
    return json.loads(r.stdout) if r.stdout.strip() else None


def round_half(x):
    return round(x * 2) / 2


def load_payload():
    text = HTML.read_text(encoding="utf-8")
    raw = re.search(r'<script id="report-data" type="application/json">(.*?)</script>', text, re.S).group(1)
    return json.loads(raw)


def load_hand():
    ws = openpyxl.load_workbook(HAND)["Ước giờ"]
    return {int(ws.cell(r, 2).value[1:]): float(ws.cell(r, 15).value) for r in range(2, 17)}


def compute(payload, hand):
    num = lambda ref: int(ref.split("#")[1])
    swp = {num(i["issueRef"]): i["finalSWP"] for i in payload["issues"]}
    parents = {num(p["issueRef"]): p for p in payload["parents"]}
    for n, p in parents.items():
        swp[n] = p["finalSWP"]
    k = sum(hand.values()) / sum(swp[n] for n in hand)
    children = {n: [num(c) for c in p["childRefs"]] for n, p in parents.items()}

    hours = {}
    for n in swp:  # issue lá: số tay nếu có, không thì SWP × k
        if n not in parents:
            hours[n] = hand.get(n, round_half(swp[n] * k))
    for n, kids in children.items():  # issue cha ước tay: chia giờ tay cho con theo SWP
        if n in hand and all(c not in parents for c in kids):
            total = sum(swp[c] for c in kids)
            for c in kids:
                hours[c] = round_half(hand[n] * swp[c] / total)

    def rollup(n):  # issue cha (kể cả epic) = tổng con; issue cha ước tay giữ số tay
        if n not in parents:
            return hours[n]
        hours[n] = hand.get(n) or sum(rollup(c) for c in children[n])
        return hours[n]
    for n in parents:
        rollup(n)
    leaf_total = sum(h for n, h in hours.items() if n not in parents)
    return swp, hours, k, leaf_total


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    swp, hours, k, leaf_total = compute(load_payload(), load_hand())
    print(f"Báo cáo: {HTML.name} · k = {k:.2f} giờ/SWP · {len(swp)} item có SWP · tổng giờ issue lá = {leaf_total:g}")

    project_id = gh("project", "view", NUMBER, "--owner", OWNER, "--format", "json")["id"]
    fields = {f["name"]: f["id"] for f in gh("project", "field-list", NUMBER, "--owner", OWNER, "--format", "json")["fields"]}
    items = gh("project", "item-list", NUMBER, "--owner", OWNER, "--format", "json", "--limit", "1000")["items"]
    changed = 0
    for it in items:
        n = it.get("content", {}).get("number")
        if n not in swp:
            continue
        for name, value in (("SWP", swp[n]), ("Giờ ước lượng", hours[n])):
            current = next((v for key, v in it.items() if key.lower() == name.lower()), None)  # gh trả khoá kiểu "sWP"
            if current is not None and float(current) == float(value):
                continue
            changed += 1
            print(f"~ #{n} {name}: {current} → {value}")
            if not a.dry_run:
                gh("project", "item-edit", "--id", it["id"], "--project-id", project_id,
                   "--field-id", fields[name], "--number", str(value))
                time.sleep(0.5)
    print(f"Ô đã đổi: {changed}" + (" (dry-run, chưa ghi)" if a.dry_run else ""))


if __name__ == "__main__":
    main()
